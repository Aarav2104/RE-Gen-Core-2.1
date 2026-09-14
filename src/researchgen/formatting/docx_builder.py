"""Builds the final downloadable research paper as a .docx file, applying the
user's chosen formatting options (page size, columns, fonts, sizes, margins)."""
from __future__ import annotations

import os
from typing import List

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from researchgen.models import ResearchMetadata, FormattingOptions, GeneratedSection
from researchgen.images.manager import ImageManager
from researchgen.citations.manager import CitationManager


def _set_font(run, name: str, size: int, bold: bool = False, italic: bool = False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def _set_page_size(section, page_size: str, margin_in: float):
    if page_size.upper() == "A4":
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    else:  # Letter
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    section.left_margin = Inches(margin_in)
    section.right_margin = Inches(margin_in)
    section.top_margin = Inches(margin_in)
    section.bottom_margin = Inches(margin_in)


def _set_columns(section, num_cols: int):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num_cols))
    cols.set(qn("w:space"), "425")  # ~0.3in gutter


def build_docx(
    metadata: ResearchMetadata,
    sections: List[GeneratedSection],
    fmt: FormattingOptions,
    images: ImageManager,
    citer: CitationManager,
    output_path: str,
) -> str:
    doc = Document()

    # ---- Title block (single column) ----
    title_section = doc.sections[0]
    _set_page_size(title_section, fmt.page_size, fmt.margin_inches)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(metadata.title or "Untitled Research Paper")
    _set_font(run, fmt.heading_font, fmt.title_size, bold=True)

    if metadata.authors:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(metadata.authors)
        _set_font(run, fmt.body_font, fmt.subheading_size)

    sub_line = ", ".join(x for x in [metadata.department, metadata.institution] if x)
    if sub_line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(sub_line)
        _set_font(run, fmt.body_font, fmt.body_size, italic=True)

    if metadata.email:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(metadata.email)
        _set_font(run, fmt.body_font, fmt.caption_size)

    if metadata.keywords:
        p = doc.add_paragraph()
        run = p.add_run("Keywords: ")
        _set_font(run, fmt.body_font, fmt.body_size, bold=True)
        run2 = p.add_run(metadata.keywords)
        _set_font(run2, fmt.body_font, fmt.body_size)

    # ---- Body: start a new section so we can switch to N columns ----
    new_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    _set_page_size(new_sec, fmt.page_size, fmt.margin_inches)
    if fmt.columns and fmt.columns > 1:
        _set_columns(new_sec, fmt.columns)

    for sec in sections:
        heading_p = doc.add_paragraph()
        run = heading_p.add_run(sec.name)
        _set_font(run, fmt.heading_font, fmt.heading_size, bold=True)

        if sec.gap_warning:
            warn_p = doc.add_paragraph()
            wrun = warn_p.add_run(f"[Note: {sec.gap_warning}]")
            _set_font(wrun, fmt.body_font, fmt.caption_size, italic=True)

        for para_text in (sec.content or "").split("\n"):
            para_text = para_text.strip()
            if not para_text:
                continue
            body_p = doc.add_paragraph()
            body_p.paragraph_format.line_spacing = fmt.line_spacing
            body_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = body_p.add_run(para_text)
            _set_font(run, fmt.body_font, fmt.body_size)

        # Insert any images assigned to this section
        for img in images.for_section(sec.name):
            if os.path.exists(img.filepath):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    img_p.add_run().add_picture(img.filepath, width=Inches(3.0))
                except Exception:
                    pass
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = cap_p.add_run(images.caption_line(img))
                _set_font(cap_run, fmt.body_font, fmt.caption_size, italic=True)

    # ---- References (single column) ----
    ref_section = doc.add_section(WD_SECTION.CONTINUOUS)
    _set_page_size(ref_section, fmt.page_size, fmt.margin_inches)
    _set_columns(ref_section, 1)

    ref_heading = doc.add_paragraph()
    run = ref_heading.add_run("References")
    _set_font(run, fmt.heading_font, fmt.heading_size, bold=True)

    for line in citer.render_references(metadata.citation_style).split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_font(run, fmt.body_font, fmt.body_size)

    doc.save(output_path)
    return output_path
